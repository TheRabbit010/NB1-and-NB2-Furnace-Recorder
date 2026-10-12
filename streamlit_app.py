import streamlit as st
import struct
import io
import re
import pandas as pd
import plotly.graph_objects as go
from datetime import datetime, timedelta
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Page Config ───────────────────────────────────────────────────────────────
st.set_page_config(
    page_title='YOKOGAWA Recorder NB1 & NB2 CAB Furnace',
    page_icon='📊',
    layout='wide',
    initial_sidebar_state='expanded',
)

# ── Styling ───────────────────────────────────────────────────────────────────
st.markdown("""
<style>
    .main-header {
        background: linear-gradient(135deg, #1F4E79, #2E75B6);
        color: white; padding: 24px 32px; border-radius: 12px;
        margin-bottom: 24px;
        box-shadow: 0 4px 6px rgba(0,0,0,0.1);
    }
    .main-header h1 { margin: 0; font-size: 1.8rem; font-weight: 600; letter-spacing: 0.5px; }
    .main-header p  { margin: 6px 0 0; opacity: 0.9; font-size: 1rem; }
    .stat-box {
        border: 1px solid var(--secondary-background-color);
        background-color: var(--secondary-background-color);
        border-radius: 10px; padding: 14px 18px; text-align: center;
    }
    .stat-box .val { font-size: 1.5rem; font-weight: 700; color: #2E75B6; }
    .stat-box .lbl { font-size: 0.78rem; opacity: 0.8; margin-top: 2px; }
</style>
""", unsafe_allow_html=True)

# ── Constants ─────────────────────────────────────────────────────────────────
ALL_CH = [
    ('CH001','Top1 (Z1)',      '°C'),
    ('CH002','Top2 (Z2)',      '°C'),
    ('CH003','Top3 (Z3)',      '°C'),
    ('CH004','Top4 (Z4)',      '°C'),
    ('CH005','Top5 (Z5)',      '°C'),
    ('CH006','Top6 (Z6)',      '°C'),
    ('CH007','Top7 (Z7)',      '°C'),
    ('CH008','Bot1 (Z8)',      '°C'),
    ('CH009','Bot2 (Z9)',      '°C'),
    ('CH010','Bot3 (Z10)',     '°C'),
    ('CH011','Bot4 (Z11)',     '°C'),
    ('CH012','Bot5 (Z12)',     '°C'),
    ('CH013','Bot6 (Z13)',     '°C'),
    ('CH014','Bot7 (Z14)',     '°C'),
    ('CH015','O2 Exit',        'ppm'),
    ('CH016','Dryer zone1',    '°C'),
    ('CH017','Dryer zone2',    '°C'),
    ('CH018','O2 Entrance',    'ppm'),
]

COLORS_TOP = ['#FF0000', '#008000', '#0000FF', '#8A2BE2', '#A52A2A', '#FFA500', '#9ACD32'] 
COLORS_BOT = ['#ADD8E6', '#FF69B4', '#808080', '#FF0000', '#008000', '#0000FF', '#8A2BE2']
COLORS_DRY = {'CH016':'#FFA500', 'CH017':'#9ACD32'} 
COLORS_O2  = {'CH015':'#A52A2A', 'CH018':'#ADD8E6'}

# ── Robust Parser ─────────────────────────────────────────────────────────────
def find_dad_params(raw):
    max_search = min(35000, len(raw) - 200)
    for offset in range(10000, max_search):
        yr, mo, dy, hr, mn, sc = raw[offset:offset+6]
        if 0 <= yr <= 99 and 1 <= mo <= 12 and 1 <= dy <= 31 and 0 <= hr <= 23 and 0 <= mn <= 59 and 0 <= sc <= 59:
            for rs in [84, 88, 92, 96, 80, 100, 128]:
                if offset + rs*2 + 6 < len(raw):
                    yr2, mo2, dy2, hr2, mn2, sc2 = raw[offset+rs : offset+rs+6]
                    yr3, mo3, dy3, hr3, mn3, sc3 = raw[offset+rs*2 : offset+rs*2+6]
                    if (0 <= yr2 <= 99 and 1 <= mo2 <= 12 and 1 <= dy2 <= 31 and 0 <= hr2 <= 23) and \
                       (0 <= yr3 <= 99 and 1 <= mo3 <= 12 and 1 <= dy3 <= 31 and 0 <= hr3 <= 23):
                        return offset, rs
    return 15008, 84

def parse_dad(raw: bytes) -> list:
    records = []
    offset, record_size = find_dad_params(raw)
    
    body_len = len(raw) - offset
    if body_len <= 0: return records
    total = body_len // record_size
    
    for i in range(total):
        base = offset + i * record_size
        hdr  = raw[base:base+8]
        if len(hdr) < 8: break
        yr, mo, dy, hr, mn, sc = hdr[0], hdr[1], hdr[2], hdr[3], hdr[4], hdr[5]
        
        if not (0 <= yr <= 99 and 1 <= mo <= 12 and 1 <= dy <= 31 and 0 <= hr <= 23 and 0 <= mn <= 59 and 0 <= sc <= 59):
            continue
            
        parsed_yr = (2000 + yr) if yr < 80 else (1900 + yr)
        try:
            ts = datetime(parsed_yr, mo, dy, hr, mn, sc)
        except ValueError:
            continue
            
        rec = {'ts': ts}
        for ci, (ch, _, _) in enumerate(ALL_CH):
            data_pos = base + 8 + ci*4
            if data_pos + 4 > len(raw): break
            
            min_v = struct.unpack_from('>h', raw, data_pos)[0]
            max_v = struct.unpack_from('>h', raw, data_pos + 2)[0]
            
            def is_valid(v):
                return v not in (-32768, -32767, 32767, 32766) and (-30000 < v < 30000)
            
            if is_valid(min_v) and is_valid(max_v):
                min_val = min_v / 10.0
                max_val = max_v / 10.0
                if -100.0 <= min_val <= 2000.0 and -100.0 <= max_val <= 2000.0:
                    rec[f'{ch}_min'] = round(min_val, 1)
                    rec[f'{ch}_max'] = round(max_val, 1)
                    rec[f'{ch}_avg'] = round((min_val + max_val) / 2.0, 1)
                    rec[ch] = rec[f'{ch}_min']
                else:
                    rec[f'{ch}_min'] = rec[f'{ch}_max'] = rec[f'{ch}_avg'] = rec[ch] = None
            else:
                rec[f'{ch}_min'] = rec[f'{ch}_max'] = rec[f'{ch}_avg'] = rec[ch] = None
                
        records.append(rec)
    return records

# ── Excel Builder ─────────────────────────────────────────────────────────────
def build_excel(data, dt_start, dt_end, zones_sel):
    hdr_fill  = PatternFill('solid', fgColor='1F4E79')
    hdr_font  = Font(bold=True, color='FFFFFF', name='Arial', size=10)
    norm_font = Font(name='Arial', size=9)
    thin      = Side(border_style='thin', color='C0C8D0')
    border    = Border(left=thin, right=thin, top=thin, bottom=thin)

    def hcell(ws, col, row, val):
        c = ws.cell(row, col, val)
        c.fill=hdr_fill; c.font=hdr_font; c.border=border
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        return c

    def dcell(ws, col, row, val, fill=None, fmt=None, bold=False):
        c = ws.cell(row, col, val)
        c.fill=fill or PatternFill('solid',fgColor='FFFFFF')
        c.font=Font(name='Arial',size=9,bold=bold); c.border=border
        c.alignment=Alignment(horizontal='right' if isinstance(val,(int,float)) else 'left')
        if fmt: c.number_format=fmt
        return c

    sel_ch = [(ch,name,unit) for ch,name,unit in ALL_CH if name in zones_sel]
    wb = Workbook()
    n  = len(data)
    report_name = f'{dt_start.strftime("%d-%m-%Y %H.%M")} to {dt_end.strftime("%d-%m-%Y %H.%M")}'

    ws = wb.active; ws.title='Raw Data'
    ws['A1'] = f'DAD Report — {report_name} | {n:,} records'
    ws.merge_cells(f'A1:{get_column_letter(3+len(sel_ch))}1')
    ws['A1'].font=Font(bold=True,name='Arial',size=12,color='1F4E79')
    ws.row_dimensions[1].height=22
    hcell(ws,1,2,'Timestamp'); hcell(ws,2,2,'Date'); hcell(ws,3,2,'Time')
    for i,(ch,name,unit) in enumerate(sel_ch):
        hcell(ws,4+i,2,f'{name}\n({ch})\n[{unit}]')
    ws.row_dimensions[2].height=48
    
    for ri,rec in enumerate(data,start=3):
        ts=rec['ts']; dp=ts.strftime('%Y/%m/%d'); tp=ts.strftime('%H:%M:%S')
        bg=PatternFill('solid',fgColor='FFFFFF')
        dcell(ws,1,ri,ts.strftime('%Y/%m/%d %H:%M:%S'),bg)
        dcell(ws,2,ri,dp,bg); dcell(ws,3,ri,tp,bg)
        for i,(ch,_,_) in enumerate(sel_ch):
            c=ws.cell(ri,4+i,rec[ch]); c.fill=bg; c.font=norm_font
            c.border=border; c.alignment=Alignment(horizontal='right'); c.number_format='0.0'
            
    ws.column_dimensions['A'].width=22; ws.column_dimensions['B'].width=13; ws.column_dimensions['C'].width=10
    for i in range(len(sel_ch)): ws.column_dimensions[get_column_letter(4+i)].width=14
    ws.freeze_panes='A3'; ws.auto_filter.ref=f'A2:{get_column_letter(3+len(sel_ch))}{n+2}'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()

# ── Chart Builder ─────────────────────────────────────────────────────────────
def render_single_chart(data, sel_names, all_channels_ref, colors_ref, title, y_range, y_title):
    if not sel_names:
        return None
        
    sel_ch_info = [(ch,n,u) for ch,n,u in all_channels_ref if n in sel_names]
    if not sel_ch_info:
        return None

    fig = go.Figure()
    ts = [r['ts'] for r in data]
    
    for i, (ch, name, unit) in enumerate(sel_ch_info):
        color = colors_ref[i] if isinstance(colors_ref, list) else colors_ref.get(ch, '#888')
        fig.add_trace(go.Scatter(
            x=ts, y=[r[ch] for r in data],
            name=name, mode='lines',
            line=dict(color=color, width=1.5),
            hovertemplate=f'<b>{name}</b>: %{{y:.1f}} {unit}<extra></extra>',
        ))

    # ปรับแต่ง Layout ให้สวยงาม สบายตา ไม่มีกรอบขาวทึบกวนใจในโหมด Dark Mode
    fig.update_layout(
        title=dict(text=title, font=dict(size=18, weight="bold")),
        height=450,
        hovermode='x unified',
        hoverlabel=dict(font_size=12, namelength=30),
        legend=dict(
            title="Channels",
            orientation="v",
            yanchor="middle", y=0.5, 
            xanchor="left", x=1.02,
            font=dict(size=12),
            itemsizing='constant'
        ),
        margin=dict(l=60, r=150, t=60, b=40),
    )
    
    fig.update_yaxes(range=y_range, title_text=y_title, title_font=dict(size=14))
    fig.update_xaxes(
        tickformat='%d/%m\n%H:%M',
        dtick=1*3600*1000, 
        tickfont=dict(size=11),
        showticklabels=True,
    )
    return fig

# ══════════════════════════════════════════════════════════════════════════════
# MAIN UI
# ══════════════════════════════════════════════════════════════════════════════

st.markdown("""
<div class="main-header">
    <h1>YOKOGAWA Recorder NB1 & NB2 CAB Furnace</h1>
    <p>Furnace Data Analysis Report</p>
</div>
""", unsafe_allow_html=True)

with st.sidebar:
    st.markdown('### 📂 Upload ไฟล์ .DAD')
    uploaded = st.file_uploader(
        'ลากไฟล์ .DAD มาวางที่นี่ (เลือกได้หลายไฟล์)',
        type=['DAD','dad'],
        accept_multiple_files=True,
    )
    
    st.markdown("---")
    st.markdown('### ⚙️ ตั้งค่าการแสดงผลกราฟ')
    value_mode = st.selectbox(
        'เลือกประเภทของค่าที่ต้องการวิเคราะห์:',
        options=['MIN', 'AVG (เฉลี่ย)', 'MAX'],
        index=2 
    )
    mode_key = {'MIN': '_min', 'AVG (เฉลี่ย)': '_avg', 'MAX': '_max'}[value_mode]

if not uploaded:
    st.info('👉 กรุณาอัปโหลดไฟล์ .DAD จากแถบเมนูด้านซ้ายเพื่อเริ่มต้นสร้างรายงานวิเคราะห์ข้อมูล')
    st.stop()

# กำหนดให้ทุกเส้นแสดงเป็นค่าเริ่มต้น
all_names_top = [name for _,name,_ in ALL_CH[:7]]
all_names_bot = [name for _,name,_ in ALL_CH[7:14]]
all_names_dry = [name for _,name,_ in ALL_CH[15:17]]
all_names_o2  = [name for _,name,_ in [ALL_CH[14], ALL_CH[17]]]

@st.cache_data(show_spinner='⏳ ระบบกำลังประมวลผลไฟล์...')
def load_data_auto(files_data, mode_k):
    all_records = []
    for fname, raw in files_data:
        recs = parse_dad(raw)
        all_records.extend(recs)
        
    if not all_records: return []
    
    all_records.sort(key=lambda x: x['ts'])
    seen = set()
    unique = []
    for r in all_records:
        k = r['ts']
        if k not in seen:
            seen.add(k)
            new_r = dict(r)
            for ch,_,_ in ALL_CH:
                new_r[ch] = r.get(f'{ch}{mode_k}')
            unique.append(new_r)
    return unique

files_data = [(f.name, f.read()) for f in uploaded]
data = load_data_auto(tuple((n,d) for n,d in files_data), mode_key)

if len(data) == 0:
    st.error('🚨 ไม่พบข้อมูลที่อ่านได้จากไฟล์ที่อัปโหลด หรือรูปแบบไฟล์ไม่ถูกต้อง')
    st.stop()

dt_start = data[0]['ts']
dt_end = data[-1]['ts']

duration = dt_end - dt_start
hrs = int(duration.total_seconds() // 3600)
mins = int((duration.total_seconds() % 3600) // 60)

col1, col2, col3, col4 = st.columns(4)
with col1:
    st.markdown(f'<div class="stat-box"><div class="val">{len(data):,}</div><div class="lbl">Total Records</div></div>', unsafe_allow_html=True)
with col2:
    st.markdown(f'<div class="stat-box"><div class="val">{len(uploaded)}</div><div class="lbl">Uploaded Files (.DAD)</div></div>', unsafe_allow_html=True)
with col3:
    st.markdown(f'<div class="stat-box"><div class="val">{hrs}h {mins}m</div><div class="lbl">Data Duration</div></div>', unsafe_allow_html=True)
with col4:
    days_with_data = len(set(r['ts'].strftime('%Y/%m/%d') for r in data))
    st.markdown(f'<div class="stat-box"><div class="val">{days_with_data}</div><div class="lbl">Active Days</div></div>', unsafe_allow_html=True)

st.markdown('<div style="margin-top:24px"></div>', unsafe_allow_html=True)

tab_chart, tab_summary, tab_download = st.tabs([
    '📈 กราฟแสดงผล (Chart)', '📊 ข้อมูลสรุป (Summary)', '⬇️ ดาวน์โหลดรายงาน (Excel)'
])

with tab_chart:
    st.markdown('<br>', unsafe_allow_html=True)
    
    # นำ MultiSelect ด้านบนกราฟออกทั้งหมด และรับค่าโซนทั้งหมดแทน
    fig_top = render_single_chart(data, all_names_top, ALL_CH[:7], COLORS_TOP, 'Top Zones Temperature', [400, 650], '°C')
    if fig_top: st.plotly_chart(fig_top, use_container_width=True, theme="streamlit")
    st.markdown('<br>', unsafe_allow_html=True)

    fig_bot = render_single_chart(data, all_names_bot, ALL_CH[7:14], COLORS_BOT, 'Bottom Zones Temperature', [400, 650], '°C')
    if fig_bot: st.plotly_chart(fig_bot, use_container_width=True, theme="streamlit")
    st.markdown('<br>', unsafe_allow_html=True)

    fig_dry = render_single_chart(data, all_names_dry, ALL_CH[15:17], COLORS_DRY, 'Dryer Zones Temperature', [150, 400], '°C')
    if fig_dry: st.plotly_chart(fig_dry, use_container_width=True, theme="streamlit")
    st.markdown('<br>', unsafe_allow_html=True)
    
    fig_o2 = render_single_chart(data, all_names_o2, [ALL_CH[14], ALL_CH[17]], COLORS_O2, 'O2 Concentration Level', [0, 150], 'ppm')
    if fig_o2: st.plotly_chart(fig_o2, use_container_width=True, theme="streamlit")

with tab_summary:
    st.markdown(f'**ช่วงเวลาของข้อมูล:** {dt_start.strftime("%d/%m/%Y %H:%M")} → {dt_end.strftime("%d/%m/%Y %H:%M")}')
    st.markdown(f'**การประมวลผลค่าที่แสดง:** {value_mode}')
    st.markdown('---')

    all_sel_names = all_names_top + all_names_bot + all_names_dry + all_names_o2
    sel_channels = [(ch,name,unit) for ch,name,unit in ALL_CH if name in all_sel_names]

    rows_summary = []
    for ch,name,unit in sel_channels:
        vals = [r[ch] for r in data if r[ch] is not None]
        if vals:
            rows_summary.append({
                'Zone / Channel': name,
                'Channel ID': ch,
                'Unit': unit,
                'Avg': round(sum(vals)/len(vals),1),
                'Min': round(min(vals),1),
                'Max': round(max(vals),1),
            })
    if rows_summary:
        df = pd.DataFrame(rows_summary)
        st.dataframe(df, use_container_width=True, hide_index=True)
    else:
        st.warning('ไม่มีข้อมูลตัวเลขในไฟล์ที่ประมวลผล')

with tab_download:
    st.markdown('### ⬇️ Export Data Report (Excel)')
    all_sel_names = all_names_top + all_names_bot + all_names_dry + all_names_o2
    
    if st.button('📊 กดเพื่อสร้างไฟล์ Excel', type='primary'):
        with st.spinner('⏳ ระบบกำลังเตรียมไฟล์ Excel...'):
            excel_data = build_excel(data, dt_start, dt_end, all_sel_names)
            fname = f'DAD_Report_NB1_NB2_{dt_start.strftime("%Y%m%d_%H%M")}.xlsx'
            
            st.download_button(
                label='⬇️ ดาวน์โหลดไฟล์ Excel (Click to Download)',
                data=excel_data,
                file_name=fname,
                mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                type='primary',
                use_container_width=True,
            )
            st.success(f'✅ ไฟล์พร้อมดาวน์โหลดแล้ว!')
